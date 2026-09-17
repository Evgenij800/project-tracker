import os
import shutil
import sqlite3
from datetime import datetime, timedelta
from io import BytesIO
from flask import Flask, render_template, request, redirect, url_for, session, g, send_from_directory, send_file
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
import docx

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'tracker-secret-key-stage-final')

DATABASE = 'tracker.db'
UPLOAD_FOLDER = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'uploads')
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

BACKUP_FOLDER = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'backups')
os.makedirs(BACKUP_FOLDER, exist_ok=True)
app.config['BACKUP_FOLDER'] = BACKUP_FOLDER

TRANSLATIONS = {
    'ru': {
        'app_title': 'Трекер Задач и Проектов',
        'login': 'Вход в систему',
        'logout': 'Выйти',
        'username': 'Имя пользователя',
        'password': 'Пароль',
        'sign_in': 'Войти',
        'test_logins': 'Тестовые учетные записи для проверки:',
        'role_admin': 'Администратор',
        'role_manager': 'Менеджер проектов',
        'role_employee': 'Сотрудник',
        'nav_dashboard': 'Главная',
        'nav_projects': 'Проекты',
        'nav_payments': 'Оплаты',
        'nav_staff': 'Кадры',
        'nav_admin_emails': 'Рассылки и SMTP',
        'nav_backups': 'Бэкапы БД',
        'welcome': 'Добро пожаловать',
        'role': 'Роль',
        'category': 'Категория',
        'projects_title': 'Модуль проектов',
        'projects_subtitle': 'Управление пулом проектов и работ по техническому заданию',
        'payments_title': 'Модуль оплат',
        'payments_subtitle': 'Мониторинг ожидающих и наступающих оплат по проектам и истории выплат',
        'staff_title': 'Модуль кадров',
        'staff_subtitle': 'База данных кадрового персонала компании и специализаций',
        'stage1_badge': 'Система управления проектами компании — Этап 6',
        'project_name': 'Название проекта',
        'description': 'Описание',
        'status': 'Статус',
        'manager': 'Менеджер',
        'actions': 'Действия',
        'add_project': 'Создать проект',
        'coming_soon': 'Функционал данного модуля будет полностью реализован в следующих итерациях согласно техническому заданию.',
        'footer': 'Система управления проектами компании',
        'profile': 'Профиль',
        'all_rights': 'Все права защищены',
        'active': 'Активен',
        'completed': 'Завершен',
        'not_assigned': 'Не назначен',
        'cancel': 'Отмена',
        'create': 'Создать',
        'return_dashboard': 'Вернуться на главную',
        'go_to_projects': 'Перейти в проекты',
        'go_to_payments': 'Перейти в оплаты',
        'go_to_staff': 'Перейти в кадры',
        'full_name': 'ФИО сотрудника',
        'login_col': 'Логин',
        'cat_admin': 'Системный администратор',
        'cat_manager': 'Проектный менеджер',
        'cat_artist': '3D Моделлер / Художник',
        'access_denied': 'Доступ запрещен. У вас недостаточно прав для просмотра этого модуля.',
        'projects_card_desc': 'Управление пулом проектов, создание работ, назначение менеджеров и исполнителей.',
        'payments_card_desc': 'Мониторинг ожидающих и наступающих оплат по проектам, история выплат исполнителям.',
        'staff_card_desc': 'База данных кадрового персонала с категоризацией (художники, моделлеры, менеджеры).',
        'active_projects_label': 'Активных:',
        'staff_count_label': 'Штат:',
        'dashboard_intro': 'Используйте навигацию выше для переключения между модулями системы:',
        'kanban_board': 'Канбан-доска проекта',
        'status_new': 'Новая',
        'status_in_progress': 'В работе',
        'status_review': 'На проверке',
        'status_completed': 'Завершена',
        'priority_low': 'Низкий',
        'priority_medium': 'Средний',
        'priority_high': 'Высокий',
        'priority': 'Приоритет',
        'assignee': 'Исполнитель',
        'due_date': 'Срок выполнения',
        'tasks': 'Задачи',
        'create_task': 'Создать задачу',
        'task_title': 'Название задачи',
        'task_description': 'Описание задачи',
        'attachments': 'Прикрепленные файлы',
        'upload_file': 'Загрузить файл',
        'comments': 'Комментарии',
        'add_comment': 'Добавить комментарий',
        'send': 'Отправить',
        'back_to_projects': 'Назад к списку проектов',
        'no_tasks': 'Нет задач',
        'task_details': 'Детали задачи',
        'files': 'Файлы',
        'open_kanban': 'Открыть канбан',
        'payments_table': 'Таблица оплат по проектам',
        'payout_history': 'История выплат и финансы',
        'upcoming_payments': 'Предстоящие оплаты (ближайшие 7 дней)',
        'amount': 'Сумма (₽)',
        'payment_date': 'Дата платежа',
        'payment_status': 'Статус',
        'status_pending': 'Ожидает',
        'status_paid': 'Оплачено',
        'staff_details': 'Детализация сотрудника',
        'assigned_projects': 'Проекты сотрудника',
        'assigned_tasks': 'Задачи сотрудника',
        'add_employee': 'Добавить сотрудника',
        'full_name_label': 'ФИО сотрудника',
        'password_label': 'Пароль',
        'category_label': 'Категория',
        'back_to_staff': 'Назад к списку кадров',
        'view_profile': 'Профиль',
        'no_upcoming_payments': 'Нет предстоящих оплат в ближайшие 7 дней.',
        'no_payments': 'Платежи не найдены.',
        'no_tasks_found': 'Задачи не найдены.',
        'no_projects_found': 'Проекты не найдены.',
        'no_comments': 'Нет комментариев.',
        'no_files': 'Нет прикрепленных файлов.',
        'email_log': 'Журнал email-уведомлений',
        'smtp_settings': 'Настройки SMTP (Исходящая почта)',
        'mass_mailing': 'Массовая рассылка штату',
        'recipient': 'Получатель',
        'subject': 'Тема письма',
        'body': 'Текст сообщения',
        'sent_date': 'Дата и время',
        'send_mailing': 'Отправить рассылку',
        'save_settings': 'Сохранить настройки',
        'project_chat': 'Чат проекта (Обсуждение)',
        'project_files': 'Общие файлы проекта',
        'upload_project_file': 'Загрузить файл',
        'delete': 'Удалить',
        'smtp_server': 'SMTP Сервер',
        'smtp_port': 'SMTP Порт',
        'smtp_user': 'SMTP Логин',
        'smtp_pass': 'SMTP Пароль',
        'backup_list': 'Резервные копии базы данных',
        'create_backup': 'Создать бэкап сейчас',
        'filename': 'Имя файла',
        'size': 'Размер',
        'created_at': 'Дата создания',
        'download': 'Скачать',
        'no_backups': 'Резервные копии не найдены.',
        'generate_contract': 'Сформировать договор',
        'generate_act': 'Сформировать акт'
    },
    'en': {
        'app_title': 'Project & Task Tracker',
        'login': 'Sign In',
        'logout': 'Logout',
        'username': 'Username',
        'password': 'Password',
        'sign_in': 'Login',
        'test_logins': 'Test credentials for each role:',
        'role_admin': 'Administrator',
        'role_manager': 'Project Manager',
        'role_employee': 'Employee',
        'nav_dashboard': 'Dashboard',
        'nav_projects': 'Projects',
        'nav_payments': 'Payments',
        'nav_staff': 'Staff',
        'nav_admin_emails': 'Mail & SMTP',
        'nav_backups': 'DB Backups',
        'welcome': 'Welcome',
        'role': 'Role',
        'category': 'Category',
        'projects_title': 'Projects Module',
        'projects_subtitle': 'Management of project pools and tasks according to technical specification',
        'payments_title': 'Payments Module',
        'payments_subtitle': 'Monitoring pending and upcoming payments by projects and payout history',
        'staff_title': 'Staff Module',
        'staff_subtitle': 'Company personnel database and specializations',
        'stage1_badge': 'Company Project Management System — Stage 6',
        'project_name': 'Project Name',
        'description': 'Description',
        'status': 'Status',
        'manager': 'Manager',
        'actions': 'Actions',
        'add_project': 'Create Project',
        'coming_soon': 'The functionality for this module will be fully implemented in subsequent iterations according to the technical specification.',
        'footer': 'Company Project Management System',
        'profile': 'Profile',
        'all_rights': 'All rights reserved',
        'active': 'Active',
        'completed': 'Completed',
        'not_assigned': 'Not assigned',
        'cancel': 'Cancel',
        'create': 'Create',
        'return_dashboard': 'Return to Dashboard',
        'go_to_projects': 'Go to Projects',
        'go_to_payments': 'Go to Payments',
        'go_to_staff': 'Go to Staff',
        'full_name': 'Full Name',
        'login_col': 'Login',
        'cat_admin': 'System Administrator',
        'cat_manager': 'Project Manager',
        'cat_artist': '3D Modeler / Artist',
        'access_denied': 'Access denied. You do not have permission to view this module.',
        'projects_card_desc': 'Management of project pools, task creation, assigning managers and executors.',
        'payments_card_desc': 'Monitoring pending and upcoming payments by projects, payout history.',
        'staff_card_desc': 'Personnel database with categorization (artists, modelers, managers).',
        'active_projects_label': 'Active:',
        'staff_count_label': 'Staff:',
        'dashboard_intro': 'Use the navigation above to switch between system modules:',
        'kanban_board': 'Project Kanban Board',
        'status_new': 'New',
        'status_in_progress': 'In Progress',
        'status_review': 'In Review',
        'status_completed': 'Completed',
        'priority_low': 'Low',
        'priority_medium': 'Medium',
        'priority_high': 'High',
        'priority': 'Priority',
        'assignee': 'Assignee',
        'due_date': 'Due Date',
        'tasks': 'Tasks',
        'create_task': 'Create Task',
        'task_title': 'Task Title',
        'task_description': 'Task Description',
        'attachments': 'Attachments',
        'upload_file': 'Upload File',
        'comments': 'Comments',
        'add_comment': 'Add Comment',
        'send': 'Send',
        'back_to_projects': 'Back to Projects',
        'no_tasks': 'No tasks',
        'task_details': 'Task Details',
        'files': 'Files',
        'open_kanban': 'Open Kanban',
        'payments_title': 'Payments Module',
        'payments_table': 'Project Payments Table',
        'payout_history': 'Payout History & Finances',
        'upcoming_payments': 'Upcoming Payments (Next 7 Days)',
        'amount': 'Amount (₽)',
        'payment_date': 'Payment Date',
        'payment_status': 'Status',
        'status_pending': 'Pending',
        'status_paid': 'Paid',
        'staff_details': 'Employee Details',
        'assigned_projects': "Employee's Projects",
        'assigned_tasks': "Employee's Tasks",
        'add_employee': 'Add Employee',
        'full_name_label': "Employee's Full Name",
        'password_label': 'Password',
        'category_label': 'Category',
        'back_to_staff': 'Back to Staff List',
        'view_profile': 'Profile',
        'no_upcoming_payments': 'No upcoming payments in the next 7 days.',
        'no_payments': 'No payments found.',
        'no_tasks_found': 'No tasks found.',
        'no_projects_found': 'No projects found.',
        'no_comments': 'No comments yet.',
        'no_files': 'No attached files.',
        'email_log': 'Email Notification Log',
        'smtp_settings': 'SMTP Settings (Outgoing Mail)',
        'mass_mailing': 'Staff Mass Mailing',
        'recipient': 'Recipient',
        'subject': 'Subject',
        'body': 'Message Body',
        'sent_date': 'Sent Date & Time',
        'send_mailing': 'Send Mailing',
        'save_settings': 'Save Settings',
        'project_chat': 'Project Discussion (Chat)',
        'project_files': 'General Project Files',
        'upload_project_file': 'Upload File',
        'delete': 'Delete',
        'smtp_server': 'SMTP Server',
        'smtp_port': 'SMTP Port',
        'smtp_user': 'SMTP User',
        'smtp_pass': 'SMTP Password',
        'backup_list': 'Database Backups',
        'create_backup': 'Create Backup Now',
        'filename': 'Filename',
        'size': 'Size',
        'created_at': 'Created At',
        'download': 'Download',
        'no_backups': 'No backups found.',
        'generate_contract': 'Generate Contract',
        'generate_act': 'Generate Act'
    }
}

CATEGORY_TRANSLATIONS = {
    'Системный администратор': {'ru': 'Системный администратор', 'en': 'System Administrator'},
    'Проектный менеджер': {'ru': 'Проектный менеджер', 'en': 'Project Manager'},
    '3D Моделлер / Художник': {'ru': '3D Моделлер / Художник', 'en': '3D Modeler / Artist'}
}

STATUS_NAMES = {
    'ru': {'new': 'Новая', 'in_progress': 'В работе', 'review': 'На проверке', 'completed': 'Завершена'},
    'en': {'new': 'New', 'in_progress': 'In Progress', 'review': 'In Review', 'completed': 'Completed'}
}

def get_db():
    db = getattr(g, '_database', None)
    if db is None:
        db = g._database = sqlite3.connect(DATABASE)
        db.row_factory = sqlite3.Row
    return db

@app.teardown_appcontext
def close_connection(exception):
    db = getattr(g, '_database', None)
    if db is not None:
        db.close()

def log_email(recipient, subject, body):
    with app.app_context():
        db = get_db()
        sent_at = datetime.now().strftime('%Y-%m-%d %H:%M')
        db.execute('INSERT INTO email_log (recipient, subject, body, sent_at) VALUES (?, ?, ?, ?)',
                   (recipient, subject, body, sent_at))
        db.commit()

def create_database_backup():
    if os.path.exists(DATABASE):
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        backup_filename = f"tracker_backup_{timestamp}.db"
        backup_path = os.path.join(BACKUP_FOLDER, backup_filename)
        shutil.copy2(DATABASE, backup_path)
        return backup_filename
    return None

def init_db():
    with app.app_context():
        db = get_db()
        cursor = db.cursor()
        
        cursor.execute('SELECT COUNT(*) FROM sqlite_master WHERE type="table"')
        has_tables = cursor.fetchone()[0] > 0
        
        if not has_tables:
            cursor.execute('''
                CREATE TABLE IF NOT EXISTS users (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    username TEXT UNIQUE NOT NULL,
                    password TEXT NOT NULL,
                    full_name TEXT NOT NULL,
                    role TEXT NOT NULL,
                    category TEXT NOT NULL
                )
            ''')
            
            cursor.execute('''
                CREATE TABLE IF NOT EXISTS projects (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    name TEXT NOT NULL,
                    description TEXT,
                    status TEXT DEFAULT 'active',
                    manager_id INTEGER,
                    FOREIGN KEY (manager_id) REFERENCES users (id)
                )
            ''')

            cursor.execute('''
                CREATE TABLE IF NOT EXISTS tasks (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    project_id INTEGER NOT NULL,
                    title TEXT NOT NULL,
                    description TEXT,
                    status TEXT DEFAULT 'new',
                    priority TEXT DEFAULT 'medium',
                    assignee_id INTEGER,
                    due_date TEXT,
                    created_at TEXT,
                    FOREIGN KEY (project_id) REFERENCES projects (id),
                    FOREIGN KEY (assignee_id) REFERENCES users (id)
                )
            ''')

            cursor.execute('''
                CREATE TABLE IF NOT EXISTS task_comments (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    task_id INTEGER NOT NULL,
                    user_id INTEGER NOT NULL,
                    comment TEXT NOT NULL,
                    created_at TEXT,
                    FOREIGN KEY (task_id) REFERENCES tasks (id),
                    FOREIGN KEY (user_id) REFERENCES users (id)
                )
            ''')

            cursor.execute('''
                CREATE TABLE IF NOT EXISTS task_files (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    task_id INTEGER NOT NULL,
                    filename TEXT NOT NULL,
                    filepath TEXT NOT NULL,
                    uploaded_at TEXT,
                    FOREIGN KEY (task_id) REFERENCES tasks (id)
                )
            ''')

            cursor.execute('''
                CREATE TABLE IF NOT EXISTS payments (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    project_id INTEGER NOT NULL,
                    amount REAL NOT NULL,
                    payment_date TEXT NOT NULL,
                    status TEXT DEFAULT 'pending',
                    description TEXT,
                    FOREIGN KEY (project_id) REFERENCES projects (id)
                )
            ''')

            cursor.execute('''
                CREATE TABLE IF NOT EXISTS email_log (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    recipient TEXT NOT NULL,
                    subject TEXT NOT NULL,
                    body TEXT NOT NULL,
                    sent_at TEXT NOT NULL
                )
            ''')

            cursor.execute('''
                CREATE TABLE IF NOT EXISTS email_settings (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    smtp_server TEXT,
                    smtp_port TEXT,
                    smtp_user TEXT,
                    smtp_password TEXT
                )
            ''')

            cursor.execute('''
                CREATE TABLE IF NOT EXISTS project_messages (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    project_id INTEGER NOT NULL,
                    user_id INTEGER NOT NULL,
                    message TEXT NOT NULL,
                    created_at TEXT NOT NULL,
                    FOREIGN KEY (project_id) REFERENCES projects (id),
                    FOREIGN KEY (user_id) REFERENCES users (id)
                )
            ''')

            cursor.execute('''
                CREATE TABLE IF NOT EXISTS project_files (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    project_id INTEGER NOT NULL,
                    filename TEXT NOT NULL,
                    filepath TEXT NOT NULL,
                    uploaded_at TEXT NOT NULL,
                    FOREIGN KEY (project_id) REFERENCES projects (id)
                )
            ''')
            
            users_data = [
                ('admin', generate_password_hash('admin123'), 'Александр Администраторов', 'admin', 'Системный администратор'),
                ('manager', generate_password_hash('manager123'), 'Марина Менеджерова', 'manager', 'Проектный менеджер'),
                ('employee', generate_password_hash('employee123'), 'Эмиль Сотрудников', 'employee', '3D Моделлер / Художник')
            ]
            cursor.executemany('''
                INSERT INTO users (username, password, full_name, role, category)
                VALUES (?, ?, ?, ?, ?)
            ''', users_data)
            
            projects_data = [
                ('Разработка 3D-персонажей для фэнтези-RPG', 'Создание детализированных игровых моделей главных героев и монстров', 'active', 2),
                ('Интеграция платежного шлюза Stripe & PayPal', 'Подключение эквайринга, настройка вебхуков и обработка транзакций', 'active', 2),
                ('Редизайн корпоративного портала студии', 'Обновление UI/UX дизайна интерфейса и адаптивной мобильной версии', 'completed', 2)
            ]
            cursor.executemany('''
                INSERT INTO projects (name, description, status, manager_id)
                VALUES (?, ?, ?, ?)
            ''', projects_data)

            now_str = datetime.now().strftime('%Y-%m-%d %H:%M')
            tasks_data = [
                (1, 'Создать высокополигональную модель дракона', 'Скульптуринг чешуи и крыльев в ZBrush по концепту', 'in_progress', 'high', 3, '2026-03-30', now_str),
                (1, 'Текстурирование доспехов палладина', 'Рисование текстур PBR в Substance Painter с картами нормалей', 'new', 'medium', 3, '2026-04-05', now_str),
                (2, 'Настройка защищенных API-эндпоинтов', 'Реализация валидации токенов и шифрования платежных данных', 'review', 'high', 2, '2026-03-25', now_str),
            ]
            cursor.executemany('''
                INSERT INTO tasks (project_id, title, description, status, priority, assignee_id, due_date, created_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            ''', tasks_data)

            comments_data = [
                (1, 1, 'Концепт утвержден арт-директором, можно приступать к детальной проработке геометрии.', now_str),
                (1, 2, 'Обратите внимание на лимит полигонов для мобильной оптимизации.', now_str)
            ]
            cursor.executemany('''
                INSERT INTO task_comments (task_id, user_id, comment, created_at)
                VALUES (?, ?, ?, ?)
            ''', comments_data)

            today = datetime.now().date()
            payments_data = [
                (1, 180000.0, (today + timedelta(days=2)).strftime('%Y-%m-%d'), 'pending', 'Аванс за разработку персонажей 3D'),
                (1, 120000.0, (today - timedelta(days=14)).strftime('%Y-%m-%d'), 'paid', 'Предоплата первого этапа'),
                (2, 250000.0, (today + timedelta(days=5)).strftime('%Y-%m-%d'), 'pending', 'Оплата интеграции платежных модулей'),
                (2, 150000.0, (today - timedelta(days=30)).strftime('%Y-%m-%d'), 'paid', 'Стартовый платеж по эквайрингу'),
                (3, 95000.0, (today + timedelta(days=15)).strftime('%Y-%m-%d'), 'pending', 'Финальный расчет за редизайн сайта')
            ]
            cursor.executemany('''
                INSERT INTO payments (project_id, amount, payment_date, status, description)
                VALUES (?, ?, ?, ?, ?)
            ''', payments_data)

            cursor.execute('INSERT INTO email_settings (smtp_server, smtp_port, smtp_user, smtp_password) VALUES ("smtp.company.com", "587", "notifications@company.com", "••••••••")')
            
            db.commit()

        # Check / create daily backup
        backups = os.listdir(BACKUP_FOLDER)
        today_str = datetime.now().strftime('%Y%m%d')
        has_today_backup = any(today_str in b for b in backups)
        if not has_today_backup:
            create_database_backup()

@app.before_request
def before_request():
    if 'lang' not in session:
        session['lang'] = 'ru'

@app.context_processor
def inject_globals():
    lang = session.get('lang', 'ru')
    def translate_cat(cat):
        if cat in CATEGORY_TRANSLATIONS:
            return CATEGORY_TRANSLATIONS[cat].get(lang, cat)
        return cat
    return {
        't': TRANSLATIONS[lang],
        'current_lang': lang,
        'translate_cat': translate_cat
    }

@app.route('/set-lang/<lang>')
def set_lang(lang):
    if lang in ['ru', 'en']:
        session['lang'] = lang
    return redirect(request.referrer or url_for('dashboard'))

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form.get('username')
        password = request.form.get('password')
        
        db = get_db()
        user = db.execute('SELECT * FROM users WHERE username = ?', (username,)).fetchone()
        
        if user and check_password_hash(user['password'], password):
            session['user_id'] = user['id']
            session['username'] = user['username']
            session['full_name'] = user['full_name']
            session['role'] = user['role']
            session['category'] = user['category']
            return redirect(url_for('dashboard'))
        else:
            error = 'Неверное имя пользователя или пароль' if session.get('lang') == 'ru' else 'Invalid username or password'
            return render_template('login.html', error=error)
            
    return render_template('login.html')

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

@app.route('/')
def dashboard():
    if 'user_id' not in session:
        return redirect(url_for('login'))
    
    db = get_db()
    projects_count = db.execute('SELECT COUNT(*) FROM projects').fetchone()[0]
    staff_count = db.execute('SELECT COUNT(*) FROM users').fetchone()[0]
    active_projects = db.execute('SELECT * FROM projects WHERE status = "active"').fetchall()
    
    return render_template('dashboard.html', 
                           projects_count=projects_count, 
                           staff_count=staff_count,
                           active_projects=active_projects)

@app.route('/projects', methods=['GET', 'POST'])
def projects():
    if 'user_id' not in session:
        return redirect(url_for('login'))
    
    db = get_db()
    if request.method == 'POST' and session.get('role') in ['admin', 'manager']:
        name = request.form.get('name')
        description = request.form.get('description')
        if name:
            db.execute('INSERT INTO projects (name, description, status, manager_id) VALUES (?, ?, "active", ?)',
                       (name, description, session['user_id']))
            db.commit()
        return redirect(url_for('projects'))
        
    projects_list = db.execute('''
        SELECT p.*, u.full_name as manager_name 
        FROM projects p 
        LEFT JOIN users u ON p.manager_id = u.id
    ''').fetchall()
    
    return render_template('projects.html', projects=projects_list)

@app.route('/projects/<int:project_id>', methods=['GET'])
def project_detail(project_id):
    if 'user_id' not in session:
        return redirect(url_for('login'))
    
    db = get_db()
    project = db.execute('''
        SELECT p.*, u.full_name as manager_name 
        FROM projects p 
        LEFT JOIN users u ON p.manager_id = u.id 
        WHERE p.id = ?
    ''', (project_id,)).fetchone()
    
    if not project:
        return redirect(url_for('projects'))
    
    tasks = db.execute('''
        SELECT t.*, u.full_name as assignee_name 
        FROM tasks t 
        LEFT JOIN users u ON t.assignee_id = u.id 
        WHERE t.project_id = ?
    ''', (project_id,)).fetchall()
    
    staff_list = db.execute('SELECT id, full_name, category FROM users').fetchall()
    
    messages = db.execute('''
        SELECT m.*, u.full_name as author_name 
        FROM project_messages m 
        JOIN users u ON m.user_id = u.id 
        WHERE m.project_id = ? 
        ORDER BY m.id ASC
    ''', (project_id,)).fetchall()
    
    files = db.execute('SELECT * FROM project_files WHERE project_id = ? ORDER BY id ASC', (project_id,)).fetchall()
    
    return render_template('project_detail.html', project=project, tasks=tasks, staff_list=staff_list, messages=messages, files=files)

@app.route('/projects/<int:project_id>/task/create', methods=['POST'])
def create_task(project_id):
    if 'user_id' not in session or session.get('role') not in ['admin', 'manager']:
        return redirect(url_for('project_detail', project_id=project_id))
    
    title = request.form.get('title')
    description = request.form.get('description')
    priority = request.form.get('priority', 'medium')
    assignee_id = request.form.get('assignee_id')
    due_date = request.form.get('due_date')
    created_at = datetime.now().strftime('%Y-%m-%d %H:%M')
    
    if title:
        db = get_db()
        cursor = db.cursor()
        cursor.execute('''
            INSERT INTO tasks (project_id, title, description, status, priority, assignee_id, due_date, created_at)
            VALUES (?, ?, ?, 'new', ?, ?, ?, ?)
        ''', (project_id, title, description, priority, assignee_id if assignee_id else None, due_date, created_at))
        db.commit()
        
        if assignee_id:
            assignee = db.execute('SELECT full_name, username FROM users WHERE id = ?', (assignee_id,)).fetchone()
            if assignee:
                lang = session.get('lang', 'ru')
                status_name = STATUS_NAMES[lang]['new']
                log_email(assignee['username'], f"Новая задача: {title}", f"Вам назначена новая задача '{title}' в проекте #{project_id}. Статус: {status_name}. Срок: {due_date or 'не указан'}.")
        
    return redirect(url_for('project_detail', project_id=project_id))

@app.route('/task/<int:task_id>/update-status', methods=['POST'])
def update_task_status(task_id):
    if 'user_id' not in session:
        return redirect(url_for('login'))
    
    new_status = request.form.get('status')
    db = get_db()
    task = db.execute('SELECT * FROM tasks WHERE id = ?', (task_id,)).fetchone()
    if task and new_status in ['new', 'in_progress', 'review', 'completed']:
        db.execute('UPDATE tasks SET status = ? WHERE id = ?', (new_status, task_id))
        db.commit()
        
        if task['assignee_id']:
            assignee = db.execute('SELECT full_name, username FROM users WHERE id = ?', (task['assignee_id'],)).fetchone()
            if assignee:
                lang = session.get('lang', 'ru')
                status_name = STATUS_NAMES[lang].get(new_status, new_status)
                log_email(assignee['username'], f"Статус задачи обновлен: {task['title']}", f"Задача '{task['title']}' переведена в статус '{status_name}'.")
                
        return redirect(url_for('project_detail', project_id=task['project_id']))
    
    return redirect(url_for('projects'))

@app.route('/projects/<int:project_id>/message', methods=['POST'])
def add_project_message(project_id):
    if 'user_id' not in session:
        return redirect(url_for('login'))
    
    msg = request.form.get('message')
    created_at = datetime.now().strftime('%Y-%m-%d %H:%M')
    if msg:
        db = get_db()
        db.execute('INSERT INTO project_messages (project_id, user_id, message, created_at) VALUES (?, ?, ?, ?)',
                   (project_id, session['user_id'], msg, created_at))
        db.commit()
    return redirect(url_for('project_detail', project_id=project_id))

@app.route('/projects/<int:project_id>/file/upload', methods=['POST'])
def upload_project_file(project_id):
    if 'user_id' not in session or session.get('role') not in ['admin', 'manager']:
        return redirect(url_for('project_detail', project_id=project_id))
        
    if 'file' in request.files:
        file = request.files['file']
        if file and file.filename:
            filename = secure_filename(file.filename)
            unique_filename = f"proj_{datetime.now().strftime('%Y%m%d%H%M%S')}_{filename}"
            filepath = os.path.join(app.config['UPLOAD_FOLDER'], unique_filename)
            file.save(filepath)
            
            uploaded_at = datetime.now().strftime('%Y-%m-%d %H:%M')
            db = get_db()
            db.execute('INSERT INTO project_files (project_id, filename, filepath, uploaded_at) VALUES (?, ?, ?, ?)',
                       (project_id, filename, unique_filename, uploaded_at))
            db.commit()
            
    return redirect(url_for('project_detail', project_id=project_id))

@app.route('/projects/<int:project_id>/file/<int:file_id>/delete', methods=['POST'])
def delete_project_file(project_id, file_id):
    if 'user_id' not in session or session.get('role') not in ['admin', 'manager']:
        return redirect(url_for('project_detail', project_id=project_id))
        
    db = get_db()
    db.execute('DELETE FROM project_files WHERE id = ? AND project_id = ?', (file_id, project_id))
    db.commit()
    return redirect(url_for('project_detail', project_id=project_id))

@app.route('/projects/<int:project_id>/contract')
def generate_contract(project_id):
    if 'user_id' not in session:
        return redirect(url_for('login'))
        
    db = get_db()
    project = db.execute('''
        SELECT p.*, u.full_name as manager_name 
        FROM projects p 
        LEFT JOIN users u ON p.manager_id = u.id 
        WHERE p.id = ?
    ''', (project_id,)).fetchone()
    
    if not project:
        return redirect(url_for('projects'))
        
    payments = db.execute('SELECT SUM(amount) as total FROM payments WHERE project_id = ?', (project_id,)).fetchone()
    total_amount = payments['total'] or 0.0
    
    doc = docx.Document()
    doc.add_heading(f"ДОГОВОР ОКАЗАНИЯ УСЛУГ № {project_id}", level=1)
    doc.add_paragraph(f"Дата составления: {datetime.now().strftime('%d.%m.%Y')}")
    doc.add_paragraph(f"Проект: {project['name']}")
    doc.add_paragraph(f"Описание: {project['description'] or '—'}")
    doc.add_paragraph(f"Руководитель проекта (Менеджер): {project['manager_name'] or 'Не назначен'}")
    doc.add_paragraph(f"Общая сумма договора: {total_amount:,.2f} руб.")
    doc.add_heading("1. Предмет договора", level=2)
    doc.add_paragraph("Исполнитель обязуется по заданию Заказчика оказать услуги по реализации проекта, указанного в настоящем договоре, а Заказчик обязуется принять и оплатить оказанные услуги.")
    doc.add_heading("2. Подписи сторон", level=2)
    doc.add_paragraph("Заказчик: _________________      Исполнитель: _________________")
    
    file_stream = BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    
    filename = f"Contract_Project_{project_id}.docx"
    return send_file(file_stream, as_attachment=True, download_name=filename, mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")

@app.route('/projects/<int:project_id>/act')
def generate_act(project_id):
    if 'user_id' not in session:
        return redirect(url_for('login'))
        
    db = get_db()
    project = db.execute('''
        SELECT p.*, u.full_name as manager_name 
        FROM projects p 
        LEFT JOIN users u ON p.manager_id = u.id 
        WHERE p.id = ?
    ''', (project_id,)).fetchone()
    
    if not project:
        return redirect(url_for('projects'))
        
    payments = db.execute('SELECT SUM(amount) as total FROM payments WHERE project_id = ?', (project_id,)).fetchone()
    total_amount = payments['total'] or 0.0
    
    doc = docx.Document()
    doc.add_heading(f"АКТ ПРИЕМА-ПЕРЕДАЧИ ВЫПОЛНЕННЫХ РАБОТ", level=1)
    doc.add_paragraph(f"по проекту: {project['name']}")
    doc.add_paragraph(f"Дата: {datetime.now().strftime('%d.%m.%Y')}")
    doc.add_paragraph(f"Менеджер проекта: {project['manager_name'] or 'Не назначен'}")
    doc.add_paragraph(f"Сумма выполненных работ: {total_amount:,.2f} руб.")
    doc.add_paragraph("Стороны подтверждают, что работы выполнены в полном объеме, надлежащим образом и в установленные сроки. Заказчик не имеет претензий по объему и качеству.")
    doc.add_heading("Подписи сторон", level=2)
    doc.add_paragraph("Заказчик: _________________      Исполнитель: _________________")
    
    file_stream = BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    
    filename = f"Act_Project_{project_id}.docx"
    return send_file(file_stream, as_attachment=True, download_name=filename, mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")

@app.route('/task/<int:task_id>/comment', methods=['POST'])
def add_task_comment(task_id):
    if 'user_id' not in session:
        return redirect(url_for('login'))
    
    comment_text = request.form.get('comment')
    created_at = datetime.now().strftime('%Y-%m-%d %H:%M')
    
    if comment_text:
        db = get_db()
        task = db.execute('SELECT project_id FROM tasks WHERE id = ?', (task_id,)).fetchone()
        if task:
            db.execute('''
                INSERT INTO task_comments (task_id, user_id, comment, created_at)
                VALUES (?, ?, ?, ?)
            ''', (task_id, session['user_id'], comment_text, created_at))
            db.commit()
            return redirect(url_for('project_detail', project_id=task['project_id']))
            
    return redirect(url_for('projects'))

@app.route('/task/<int:task_id>/upload', methods=['POST'])
def upload_task_file(task_id):
    if 'user_id' not in session:
        return redirect(url_for('login'))
    
    db = get_db()
    task = db.execute('SELECT project_id FROM tasks WHERE id = ?', (task_id,)).fetchone()
    if not task:
        return redirect(url_for('projects'))
        
    if 'file' in request.files:
        file = request.files['file']
        if file and file.filename:
            filename = secure_filename(file.filename)
            unique_filename = f"{datetime.now().strftime('%Y%m%d%H%M%S')}_{filename}"
            filepath = os.path.join(app.config['UPLOAD_FOLDER'], unique_filename)
            file.save(filepath)
            
            uploaded_at = datetime.now().strftime('%Y-%m-%d %H:%M')
            db.execute('''
                INSERT INTO task_files (task_id, filename, filepath, uploaded_at)
                VALUES (?, ?, ?, ?)
            ''', (task_id, filename, unique_filename, uploaded_at))
            db.commit()
            
    return redirect(url_for('project_detail', project_id=task['project_id']))

@app.route('/task/<int:task_id>/data', methods=['GET'])
def get_task_data(task_id):
    if 'user_id' not in session:
        return {'error': 'Unauthorized'}, 401
        
    db = get_db()
    task = db.execute('''
        SELECT t.*, u.full_name as assignee_name 
        FROM tasks t 
        LEFT JOIN users u ON t.assignee_id = u.id 
        WHERE t.id = ?
    ''', (task_id,)).fetchone()
    
    if not task:
        return {'error': 'Not found'}, 404
        
    comments = db.execute('''
        SELECT c.*, u.full_name as author_name 
        FROM task_comments c 
        JOIN users u ON c.user_id = u.id 
        WHERE c.task_id = ? 
        ORDER BY c.id ASC
    ''', (task_id,)).fetchall()
    
    files = db.execute('SELECT * FROM task_files WHERE task_id = ? ORDER BY id ASC', (task_id,)).fetchall()
    
    return {
        'id': task['id'],
        'title': task['title'],
        'description': task['description'],
        'status': task['status'],
        'priority': task['priority'],
        'assignee': task['assignee_name'] or 'Не назначен',
        'due_date': task['due_date'] or '',
        'comments': [{'author': c['author_name'], 'text': c['comment'], 'date': c['created_at']} for c in comments],
        'files': [{'id': f['id'], 'name': f['filename'], 'url': f['filepath'], 'date': f['uploaded_at']} for f in files]
    }

@app.route('/uploads/<filename>')
def uploaded_file(filename):
    if 'user_id' not in session:
        return redirect(url_for('login'))
    return send_from_directory(app.config['UPLOAD_FOLDER'], filename)

@app.route('/payments')
def payments():
    if 'user_id' not in session:
        return redirect(url_for('login'))
    
    if session.get('role') == 'employee':
        return redirect(url_for('dashboard'))
    
    db = get_db()
    payments_list = db.execute('''
        SELECT pay.*, p.name as project_name 
        FROM payments pay 
        JOIN projects p ON pay.project_id = p.id 
        ORDER BY pay.payment_date ASC
    ''').fetchall()
    
    today = datetime.now().date()
    upcoming_limit = today + timedelta(days=7)
    
    return render_template('payments.html', payments=payments_list, today=today, upcoming_limit=upcoming_limit)

@app.route('/staff')
def staff():
    if 'user_id' not in session:
        return redirect(url_for('login'))
    
    if session.get('role') == 'employee':
        return redirect(url_for('dashboard'))
    
    db = get_db()
    staff_list = db.execute('SELECT id, username, full_name, role, category FROM users').fetchall()
    return render_template('staff.html', staff=staff_list)

@app.route('/staff/create', methods=['POST'])
def create_staff():
    if 'user_id' not in session or session.get('role') != 'admin':
        return redirect(url_for('dashboard'))
        
    username = request.form.get('username')
    password = request.form.get('password')
    full_name = request.form.get('full_name')
    role = request.form.get('role', 'employee')
    category = request.form.get('category', '3D Моделлер / Художник')
    
    if username and password and full_name:
        db = get_db()
        try:
            db.execute('''
                INSERT INTO users (username, password, full_name, role, category)
                VALUES (?, ?, ?, ?, ?)
            ''', (username, generate_password_hash(password), full_name, role, category))
            db.commit()
        except sqlite3.IntegrityError:
            pass
            
    return redirect(url_for('staff'))

@app.route('/staff/<int:user_id>')
def staff_detail(user_id):
    if 'user_id' not in session:
        return redirect(url_for('login'))
    
    if session.get('role') == 'employee' and session.get('user_id') != user_id:
        return redirect(url_for('dashboard'))
        
    db = get_db()
    employee = db.execute('SELECT id, username, full_name, role, category FROM users WHERE id = ?', (user_id,)).fetchone()
    if not employee:
        return redirect(url_for('staff'))
        
    projects_managed = db.execute('SELECT * FROM projects WHERE manager_id = ?', (user_id,)).fetchall()
    tasks_assigned = db.execute('''
        SELECT t.*, p.name as project_name 
        FROM tasks t 
        JOIN projects p ON t.project_id = p.id 
        WHERE t.assignee_id = ?
    ''', (user_id,)).fetchall()
    
    return render_template('staff_detail.html', employee=employee, projects_managed=projects_managed, tasks_assigned=tasks_assigned)

@app.route('/admin/emails', methods=['GET', 'POST'])
def admin_emails():
    if 'user_id' not in session or session.get('role') != 'admin':
        return redirect(url_for('dashboard'))
        
    db = get_db()
    
    if request.method == 'POST':
        action = request.form.get('action')
        if action == 'save_settings':
            server = request.form.get('smtp_server')
            port = request.form.get('smtp_port')
            user = request.form.get('smtp_user')
            password = request.form.get('smtp_password')
            db.execute('UPDATE email_settings SET smtp_server = ?, smtp_port = ?, smtp_user = ?, smtp_password = ? WHERE id = 1',
                       (server, port, user, password))
            db.commit()
        elif action == 'send_mailing':
            subject = request.form.get('subject')
            body = request.form.get('body')
            if subject and body:
                staff_members = db.execute('SELECT username FROM users').fetchall()
                for member in staff_members:
                    log_email(member['username'], subject, body)
        return redirect(url_for('admin_emails'))
        
    settings = db.execute('SELECT * FROM email_settings WHERE id = 1').fetchone()
    email_logs = db.execute('SELECT * FROM email_log ORDER BY id DESC').fetchall()
    
    return render_template('admin_emails.html', settings=settings, email_logs=email_logs)

@app.route('/admin/backups', methods=['GET', 'POST'])
def admin_backups():
    if 'user_id' not in session or session.get('role') != 'admin':
        return redirect(url_for('dashboard'))
        
    if request.method == 'POST':
        create_database_backup()
        return redirect(url_for('admin_backups'))
        
    backups = []
    for filename in os.listdir(BACKUP_FOLDER):
        if filename.endswith('.db'):
            path = os.path.join(BACKUP_FOLDER, filename)
            size = os.path.getsize(path)
            created_at = datetime.fromtimestamp(os.path.getmtime(path)).strftime('%Y-%m-%d %H:%M')
            backups.append({'filename': filename, 'size': f"{size / 1024:.1f} KB", 'created_at': created_at})
            
    backups.sort(key=lambda x: x['filename'], reverse=True)
    return render_template('backups.html', backups=backups)

@app.route('/admin/backups/download/<filename>')
def download_backup(filename):
    if 'user_id' not in session or session.get('role') != 'admin':
        return redirect(url_for('dashboard'))
    return send_from_directory(BACKUP_FOLDER, filename, as_attachment=True)

if __name__ == '__main__':
    init_db()
    port = int(os.environ.get('PORT', 5000))
    debug_mode = os.environ.get('FLASK_DEBUG') == '1' or os.environ.get('PORT') is None
    app.run(debug=debug_mode, host='0.0.0.0', port=port)
